"""Обработка документа: полный проход от исходника до отчёта."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

from fixtures.rich_docx import build_rich_docx
from word_formatter.core.errors import SaveError
from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.processing import (
    ProcessingResult,
    default_output_name,
    ensure_output_is_safe,
    process_document,
)


@pytest.fixture()
def profile(profile_file: Path):
    return load_profile_from_path(profile_file)


def test_output_name_is_derived_from_source() -> None:
    assert default_output_name("Курсовая работа.docx") == "Курсовая работа_оформлено.docx"
    assert default_output_name(Path("отчёт.DOCX")) == "отчёт_оформлено.DOCX"
    assert default_output_name("без расширения") == "без расширения_оформлено.docx"


def test_saving_over_source_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "работа.docx"
    source.write_bytes(b"PK\x03\x04")

    with pytest.raises(SaveError) as error:
        ensure_output_is_safe(source, source)

    assert "поверх исходного" in error.value.user_message


def test_same_path_written_differently_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "работа.docx"
    source.write_bytes(b"PK\x03\x04")

    with pytest.raises(SaveError):
        ensure_output_is_safe(source, tmp_path / "." / "работа.docx")


def test_processing_produces_separate_file(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")
    digest_before = hashlib.sha256(source.read_bytes()).hexdigest()
    target = tmp_path / "результат.docx"

    result = process_document(source, target, profile)

    assert isinstance(result, ProcessingResult)
    assert target.exists()
    assert target != source
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest_before


def test_processing_verifies_content(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")
    result = process_document(source, tmp_path / "результат.docx", profile)

    assert result.verification.passed is True, result.verification.problems
    assert result.verification.notes


def test_processing_builds_report(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")
    result = process_document(source, tmp_path / "результат.docx", profile)

    text = result.report.as_text()
    assert "Отчёт об оформлении документа" in text
    assert "исходник.docx" in text
    assert "Проверка сохранности содержимого" in text
    assert result.report.manual_checks


def test_processing_reports_changes(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")
    result = process_document(source, tmp_path / "результат.docx", profile)

    categories = result.changes.category_counts()
    assert categories["page"] == 1
    assert categories["body-text"] == 1
    assert categories["blank-lines"] >= 1


def test_processing_keeps_no_temporary_files(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")
    process_document(source, tmp_path / "результат.docx", profile)

    leftovers = [path.name for path in tmp_path.iterdir() if path.name.startswith(".word-formatter-")]
    assert leftovers == []


def test_missing_target_directory_is_reported(tmp_path: Path, profile) -> None:
    source = build_rich_docx(tmp_path / "исходник.docx")

    with pytest.raises(SaveError) as error:
        process_document(source, tmp_path / "нет-папки" / "результат.docx", profile)

    assert "не найдена" in error.value.user_message


def test_result_matches_profile(tmp_path: Path, profile) -> None:
    from docx import Document
    from docx.shared import Pt

    source = build_rich_docx(tmp_path / "исходник.docx")
    result = process_document(source, tmp_path / "результат.docx", profile)

    document = Document(str(result.output_path))
    paragraph = next(p for p in document.paragraphs if p.text.strip().startswith("Текст заключения"))
    assert paragraph.runs[0].font.size == Pt(14)

    table = document.tables[0]
    assert table.cell(1, 1).paragraphs[0].runs[0].font.size == Pt(12)
