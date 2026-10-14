"""Главный контракт приложения на настоящем документе.

Проверяется утверждение «меняем оформление, а не содержание»: после работы
правил текст, сноски, изображения, поля и связи обязаны остаться прежними,
а измениться могут только свойства оформления.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from fixtures.rich_docx import BYTE_IDENTICAL_PARTS
from support import (
    effective_content_types,
    non_formatting_difference,
    read_parts,
    relationships,
    sha256,
    texts,
)
from word_formatter.core.docx import inspect_document, open_document
from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import RuleContext, apply_all_rules

TEXT_PARTS = [
    "word/document.xml",
    "word/footnotes.xml",
    "word/comments.xml",
    "word/header1.xml",
    "word/footer1.xml",
]

BODY_PARAGRAPH = "Текст заключения во втором разделе документа."


@pytest.fixture()
def processed(rich_docx: Path, profile_file: Path, tmp_path: Path) -> SimpleNamespace:
    digest_before = sha256(rich_docx.read_bytes())

    document = open_document(rich_docx)
    profile = load_profile_from_path(profile_file)
    context = RuleContext(
        document=document,
        profile=profile,
        inventory=inspect_document(document, profile),
    )
    apply_all_rules(context)

    target = tmp_path / "оформлено.docx"
    document.save(str(target))

    return SimpleNamespace(
        source=rich_docx,
        target=target,
        context=context,
        digest_before=digest_before,
        before=read_parts(rich_docx),
        after=read_parts(target),
    )


def test_text_is_unchanged(processed: SimpleNamespace) -> None:
    for name in TEXT_PARTS:
        assert texts(processed.before[name]) == texts(processed.after[name]), f"изменился текст в {name}"


def test_only_formatting_changed(processed: SimpleNamespace) -> None:
    """Правило свободных строк вправе добавлять только пустые абзацы.

    Что добавленный абзац пуст, проверяет отдельный тест ниже: сравнение
    текста и состава элементов не даёт протащить в документ содержание.
    """
    difference = non_formatting_difference(
        processed.before["word/document.xml"],
        processed.after["word/document.xml"],
        allowed_additions=frozenset({"p"}),
    )
    assert difference == [], f"затронуто не только оформление: {difference[:3]}"


def test_only_empty_paragraphs_were_added(processed: SimpleNamespace) -> None:
    """Все абзацы результата с текстом совпадают с исходными, и их столько же."""
    source = Document(str(processed.source))
    result = Document(str(processed.target))

    source_texts = [p.text for p in source.paragraphs if p.text.strip()]
    result_texts = [p.text for p in result.paragraphs if p.text.strip()]

    assert result_texts == source_texts
    assert len(result.paragraphs) > len(source.paragraphs)
    assert result_texts, "в документе должен остаться текст"


def test_hard_parts_are_byte_identical(processed: SimpleNamespace) -> None:
    changed = [
        name
        for name in BYTE_IDENTICAL_PARTS
        if sha256(processed.before[name]) != sha256(processed.after[name])
    ]
    assert changed == []


def test_content_types_and_links_are_unchanged(processed: SimpleNamespace) -> None:
    names = sorted(set(processed.before) | set(processed.after))
    assert effective_content_types(
        processed.before["[Content_Types].xml"], names
    ) == effective_content_types(processed.after["[Content_Types].xml"], names)
    assert relationships(processed.before["word/_rels/document.xml.rels"]) == relationships(
        processed.after["word/_rels/document.xml.rels"]
    )


def test_source_file_is_untouched(processed: SimpleNamespace) -> None:
    assert sha256(processed.source.read_bytes()) == processed.digest_before


def test_body_paragraph_receives_profile_formatting(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    paragraph = next(p for p in document.paragraphs if p.text.strip() == BODY_PARAGRAPH)

    assert paragraph.paragraph_format.line_spacing == 1.5
    assert paragraph.paragraph_format.space_after == Pt(0)
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(1.25).twips

    runs = [run for run in paragraph.runs if run.text.strip()]
    assert runs, "в абзаце должен остаться текст"
    for run in runs:
        assert run.font.name == "Times New Roman"
        assert run.font.size == Pt(14)


def test_emphasis_is_removed_and_reported(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    paragraph = next(p for p in document.paragraphs if "курсивом" in p.text)
    for run in paragraph.runs:
        assert run.italic is not True

    assert processed.context.changes.count("italic-removed") >= 1
    assert processed.context.changes.count("underline-removed") >= 1
    reported = " ".join(
        change.text
        for change in processed.context.changes.items
        if change.category in {"italic-removed", "underline-removed"}
    )
    assert "Обычный текст с ручным шрифтом" in reported


def test_footnote_reference_keeps_its_style(processed: SimpleNamespace) -> None:
    document_xml = processed.after["word/document.xml"].decode("utf-8")
    assert "FootnoteReference" in document_xml
    assert "footnoteReference" in document_xml


def test_page_settings_are_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    for section in document.sections:
        assert section.left_margin.twips == Mm(30).twips
        assert section.page_width.twips == Mm(210).twips


def test_changes_are_reported_by_category(processed: SimpleNamespace) -> None:
    categories = processed.context.changes.category_counts()
    assert "page" in categories
    assert "document-defaults" in categories
    assert categories["body-text"] == 1


def test_normal_style_is_updated_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    assert document.styles["Normal"].font.name == "Times New Roman"

    rPr = (
        document.styles.element.find(qn("w:docDefaults"))
        .find(qn("w:rPrDefault"))
        .find(qn("w:rPr"))
    )
    assert rPr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Times New Roman"


def test_result_reopens_without_errors(processed: SimpleNamespace) -> None:
    """Результат должен открываться заново и содержать все части документа."""
    source = Document(str(processed.source))
    document = Document(str(processed.target))

    assert len(document.paragraphs) >= len(source.paragraphs)
    assert len(document.tables) == 1
    assert len(document.sections) == 2
    assert len(document.inline_shapes) == 1


def test_page_number_is_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    footer = document.sections[0].footer
    instructions = [
        node.text
        for paragraph in footer.paragraphs
        for node in paragraph._p.iter(qn("w:instrText"))
    ]
    assert any("PAGE" in (text or "").upper() for text in instructions)
    assert document.sections[0].different_first_page_header_footer is True


def test_blank_lines_are_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    texts_list = [paragraph.text for paragraph in document.paragraphs]
    heading_index = texts_list.index("ВВЕДЕНИЕ")
    assert texts_list[heading_index + 1] == "", "после заголовка должна идти свободная строка"

    caption_index = next(
        index for index, text in enumerate(texts_list) if text.startswith("Таблица 1")
    )
    assert texts_list[caption_index - 1] == "", "перед названием таблицы должна быть свободная строка"


def test_structural_heading_is_formatted_and_text_kept(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    paragraph = next(p for p in document.paragraphs if p.text.strip() == "ВВЕДЕНИЕ")

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.runs[0].font.bold is True
    assert paragraph.runs[0].font.all_caps is True
    assert paragraph.text.strip() == "ВВЕДЕНИЕ"


def test_figure_caption_is_formatted(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    paragraph = next(p for p in document.paragraphs if p.text.strip().startswith("Рисунок 1"))

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.runs[0].font.size == Pt(14)


def test_table_cells_are_formatted_by_content(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    table = document.tables[0]

    assert table.cell(0, 0).paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.cell(1, 0).paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.cell(1, 1).paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.cell(1, 1).paragraphs[0].runs[0].font.size == Pt(12)
    # Ячейка-прочерк не трогалась
    assert table.cell(2, 1).paragraphs[0].paragraph_format.alignment is None


def test_hyphenation_is_enabled_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    auto = document.settings.element.find(qn("w:autoHyphenation"))
    assert auto is not None
    assert auto.get(qn("w:val")) == "true"


def test_footnote_style_is_updated_in_result(processed: SimpleNamespace) -> None:
    document = Document(str(processed.target))
    style = next(style for style in document.styles if (style.style_id or "") == "FootnoteText")
    assert style.font.size == Pt(10)
    assert style.font.name == "Times New Roman"


def test_new_rule_categories_appear_in_report(processed: SimpleNamespace) -> None:
    categories = processed.context.changes.category_counts()
    for expected in (
        "structural-headings",
        "headings",
        "heading-styles",
        "figure-captions",
        "table-captions",
        "tables",
        "hyphenation",
        "footnotes",
    ):
        assert expected in categories, f"в отчёте нет раздела {expected}"
