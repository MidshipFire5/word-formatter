"""Обработка примера документа: проверка на «наивно оформленной» курсовой.

Этот тест появился после визуальной приёмки: рендеринг примера показал два
дефекта, которых не видели тесты на синтетическом эталоне — синие заголовки
из встроенного стиля Word и строку содержания с отточием, принятую
за заголовок. Тест закрепляет исправления.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from tools.make_example import build_example
from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.processing import process_document


@pytest.fixture()
def processed(tmp_path: Path, profile_file: Path):
    source = build_example(tmp_path / "пример-курсовой.docx")
    digest_before = hashlib.sha256(source.read_bytes()).hexdigest()
    profile = load_profile_from_path(profile_file)
    result = process_document(source, tmp_path / "пример-курсовой_оформлено.docx", profile)
    return source, digest_before, result


def test_source_file_is_untouched(processed) -> None:
    source, digest_before, _result = processed
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest_before


def test_content_is_preserved(processed) -> None:
    _source, _digest, result = processed
    assert result.verification.passed is True, result.verification.problems


def test_body_text_gets_profile_formatting(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    paragraph = next(p for p in document.paragraphs if p.text.startswith("Актуальность темы"))
    assert paragraph.paragraph_format.line_spacing == 1.5
    assert paragraph.runs[0].font.size == Pt(14)
    assert paragraph.runs[0].font.name == "Times New Roman"


def test_emphasis_is_removed_everywhere(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            assert run.italic is not True, f"остался курсив: {run.text!r}"
            assert not run.underline, f"осталось подчёркивание: {run.text!r}"


def test_author_colour_is_kept(processed) -> None:
    """Профиль не задаёт цвет основного текста — цвет автора не трогаем."""
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    paragraph = next(p for p in document.paragraphs if "Arial" in p.text)
    colour = paragraph.runs[0]._element.rPr.find(qn("w:color"))
    assert colour is not None
    assert colour.get(qn("w:val")) == "1F3FBF"


def test_headings_are_not_blue(processed) -> None:
    """Встроенный синий цвет стилей заголовков должен быть снят."""
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    for level in (1, 2):
        rPr = document.styles[f"Heading {level}"].element.find(qn("w:rPr"))
        assert rPr.find(qn("w:color")).get(qn("w:val")) == "auto"


def test_contents_line_did_not_become_a_heading(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    contents = [
        paragraph
        for paragraph in document.paragraphs
        if "…" in paragraph.text or "...." in paragraph.text
    ]
    assert contents, "в примере должны быть строки содержания"
    for paragraph in contents:
        assert not paragraph.style.name.startswith("Heading"), (
            f"строка содержания получила стиль заголовка: {paragraph.text!r}"
        )


def test_real_headings_are_styled(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    styles = {
        paragraph.text.strip(): paragraph.style.name
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    }
    assert styles["ВВЕДЕНИЕ"] == "Heading 1"
    assert styles["ЗАКЛЮЧЕНИЕ"] == "Heading 1"
    assert styles["1.1 Понятие и сущность процесса"] == "Heading 2"


def test_blank_lines_are_inserted(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))
    texts = [paragraph.text for paragraph in document.paragraphs]

    index = texts.index("ВВЕДЕНИЕ")
    assert texts[index + 1] == ""


def test_table_is_formatted(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))
    table = document.tables[0]

    assert table.cell(1, 1).paragraphs[0].runs[0].font.size == Pt(12)
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assert table.cell(1, 1).paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.cell(2, 1).paragraphs[0].paragraph_format.alignment is None  # прочерк


def test_page_number_is_added(processed) -> None:
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    instructions = [
        node.text
        for paragraph in document.sections[0].footer.paragraphs
        for node in paragraph._p.iter(qn("w:instrText"))
    ]
    assert any("PAGE" in (text or "").upper() for text in instructions)


def test_title_page_keeps_its_layout(processed) -> None:
    """Титульный лист методичка не описывает: его вёрстку не трогаем."""
    _source, _digest, result = processed
    document = Document(str(result.output_path))

    from docx.enum.text import WD_ALIGN_PARAGRAPH

    titles = [p for p in document.paragraphs if p.text.startswith("КРАСНОДАРСКИЙ ФИЛИАЛ")]
    assert titles
    for paragraph in titles:
        assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert paragraph.runs[0].font.size == Pt(12)


def test_report_explains_title_page(processed) -> None:
    _source, _digest, result = processed
    text = result.report.as_text()

    assert "Титульный лист" in text
    assert "оставлен без изменений" in text


def test_figure_caption_is_placed_below_the_figure(processed) -> None:
    """В примере подпись стоит над рисунком — приложение переносит её под рисунок."""
    source, _digest, result = processed

    def caption_and_figure(path: Path) -> tuple[int, int]:
        document = Document(str(path))
        caption = next(
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip().startswith("Рисунок 1")
        )
        figure = next(
            index
            for index, paragraph in enumerate(document.paragraphs)
            if next(iter(paragraph._p.iter(qn("w:drawing"))), None) is not None
        )
        return caption, figure

    source_caption, source_figure = caption_and_figure(source)
    assert source_caption < source_figure, "в примере подпись специально поставлена над рисунком"

    result_caption, result_figure = caption_and_figure(result.output_path)
    assert result_caption > result_figure, "после обработки подпись должна стоять под рисунком"


def test_verification_does_not_alarm_on_caption_move(processed) -> None:
    """Перенос подписи не считается потерей: проверка сохранности проходит."""
    _source, _digest, result = processed

    assert result.verification.passed is True, result.verification.problems


def test_report_mentions_caption_move(processed) -> None:
    _source, _digest, result = processed
    text = result.report.as_text()

    assert "перенесена под рисунок" in text
    assert "Положение подписей" in text


def test_report_lists_manual_checks(processed) -> None:
    _source, _digest, result = processed
    manual = "\n".join(result.report.manual_checks)

    assert "титульный лист" in manual.lower()
    assert result.report.as_text().count("вставлена свободная строка") > 0
