"""Правила оформления страницы."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_page_settings


def test_page_settings_match_profile(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_settings(context)

    section = document.sections[0]
    # Размеры страницы и поля хранятся в твипах: сравниваем точные значения.
    assert section.page_width.twips == Mm(210).twips
    assert section.page_height.twips == Mm(297).twips
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert section.left_margin.twips == Mm(30).twips
    assert section.right_margin.twips == Mm(15).twips
    assert section.top_margin.twips == Mm(20).twips
    assert section.bottom_margin.twips == Mm(20).twips
    assert section.header_distance.twips == Mm(20).twips
    assert section.footer_distance.twips == Mm(12.5).twips


def test_change_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_settings(context)

    assert context.changes.count("page") == 1
    change = context.changes.by_category()["page"][0]
    assert "A4" in change.description
    assert "30/15/20/20" in change.description
    assert "разделов документа: 1" == change.detail


def test_every_section_is_updated(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_section()
    document.add_section()
    assert len(document.sections) == 3

    context = make_context(document, load_profile_from_path(profile_file))
    apply_page_settings(context)

    for section in document.sections:
        assert section.left_margin.twips == Mm(30).twips
        assert section.page_width.twips == Mm(210).twips
    assert "разделов документа: 3" in context.changes.by_category()["page"][0].detail


def test_landscape_profile_swaps_dimensions(make_context, profile_dict: dict, tmp_path: Path) -> None:
    landscape = copy.deepcopy(profile_dict)
    landscape["page"]["orientation"] = "landscape"
    path = tmp_path / "landscape.json"
    path.write_text(json.dumps(landscape, ensure_ascii=False), encoding="utf-8")

    document = Document()
    context = make_context(document, load_profile_from_path(path))
    apply_page_settings(context)

    section = document.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width.twips == Mm(297).twips
    assert section.page_height.twips == Mm(210).twips


def test_page_size_is_written_to_xml(make_context, profile_file: Path) -> None:
    """Размер листа должен попасть в XML, а не остаться только в модели."""
    from docx.oxml.ns import qn

    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))
    apply_page_settings(context)

    pgSz = document.sections[0]._sectPr.find(qn("w:pgSz"))
    assert pgSz.get(qn("w:w")) == str(Mm(210).twips)
    assert pgSz.get(qn("w:h")) == str(Mm(297).twips)
