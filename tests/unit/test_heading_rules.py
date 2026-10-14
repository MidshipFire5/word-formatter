"""Правила оформления заголовков."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_heading_styles, apply_headings


def test_structural_heading_is_formatted(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Введение")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert paragraph.paragraph_format.line_spacing == 1.0

    run = paragraph.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)
    assert run.font.bold is True
    # Прописные включаются оформлением, а не правкой текста
    assert run.font.all_caps is True
    assert paragraph.text == "Введение"


def test_structural_heading_gets_word_style(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("ЗАКЛЮЧЕНИЕ")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    assert paragraph.style.name == "Heading 1"


def test_structural_heading_forbids_hyphenation(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    pPr = paragraph._p.pPr
    node = pPr.find(qn("w:suppressAutoHyphenation"))
    assert node is not None
    assert node.get(qn("w:val")) == "true"


def test_section_heading_is_formatted_by_level(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("1.1 Понятие и сущность процесса")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(1.25).twips
    assert paragraph.paragraph_format.line_spacing == 1.0
    assert paragraph.runs[0].font.bold is True
    assert paragraph.style.name == "Heading 2"


def test_numbered_heading_stays_uppercase_free(make_context, profile_file: Path) -> None:
    """Разделы печатаются как обычно: прописные — только у структурных частей."""
    document = Document()
    paragraph = document.add_paragraph("1 Теоретические основы")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    assert paragraph.runs[0].font.all_caps in (None, False)


def test_heading_style_definitions_are_updated(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_heading_styles(context)

    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        assert style.font.name == "Times New Roman"
        assert style.font.size == Pt(14)
        assert style.font.bold is True
        assert style.paragraph_format.line_spacing == 1.0
        assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert style.paragraph_format.first_line_indent.twips == Cm(1.25).twips


def test_heading_style_colour_becomes_automatic(make_context, profile_file: Path) -> None:
    """Встроенные стили заголовков Word синие.

    Приложение назначает эти стили абзацам, поэтому цвет нужно задать явно:
    иначе документ получит синие заголовки, которых у автора не было.
    """
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_heading_styles(context)

    for level in (1, 2, 3):
        rPr = document.styles[f"Heading {level}"].element.find(qn("w:rPr"))
        colour = rPr.find(qn("w:color"))
        assert colour is not None, f"у стиля заголовка {level} цвет не задан"
        assert colour.get(qn("w:val")) == "auto"


def test_profile_colour_is_used_when_specified(
    make_context, profile_dict: dict, tmp_path: Path
) -> None:
    data = copy.deepcopy(profile_dict)
    data["headings"]["levels"]["1"]["color"] = "#112233"
    path = tmp_path / "цветной.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    context = make_context(document, load_profile_from_path(path))
    apply_heading_styles(context)

    rPr = document.styles["Heading 1"].element.find(qn("w:rPr"))
    assert rPr.find(qn("w:color")).get(qn("w:val")) == "112233"


def test_heading_style_change_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_heading_styles(context)

    assert context.changes.count("heading-styles") == 1
    assert "обновлено стилей заголовков: 3" in context.changes.by_category()["heading-styles"][0].detail


def test_already_styled_heading_is_not_reassigned(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Название раздела", style="Heading 1")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    details = [change.detail for change in context.changes.items]
    assert not any("стиль заголовка" in detail for detail in details)


def test_headings_are_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("1.1 Первый подраздел")
    document.add_paragraph("1.2 Второй подраздел")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    assert context.changes.count("structural-headings") == 1
    level_change = context.changes.by_category()["headings"][0]
    assert "заголовков уровня 2: 2" in level_change.detail
    structural = context.changes.by_category()["structural-headings"][0]
    assert structural.text == "ВВЕДЕНИЕ"
    assert structural.paragraph_index == 0


def test_document_without_headings_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Просто текст без заголовков.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_headings(context)

    messages = [change.description for change in context.changes.by_category()["headings"]]
    assert any("не найдены" in message for message in messages)
