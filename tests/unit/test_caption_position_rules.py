"""Положение подписей рисунков и названий таблиц.

Методичка требует: подпись рисунка — под рисунком, название таблицы — над
таблицей. Это единственное правило, которое меняет порядок элементов,
поэтому проверяется особенно внимательно.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_caption_positions

DRAWING = (
    f'<w:drawing {nsdecls("w")}><wp:inline {nsdecls("wp")}>'
    f'<wp:extent cx="100" cy="100"/></wp:inline></w:drawing>'
)


def add_figure(document: Document, caption: str | None = None, *, caption_first: bool = False):
    """Добавляет рисунок и, если нужно, подпись до или после него."""
    if caption_first and caption:
        document.add_paragraph(caption)
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml(DRAWING))
    if not caption_first and caption:
        document.add_paragraph(caption)
    return paragraph


def add_table(document: Document, caption: str | None = None, *, caption_first: bool = False):
    """Добавляет таблицу и, если нужно, название до или после неё."""
    if caption_first and caption:
        document.add_paragraph(caption)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Данные"
    if not caption_first and caption:
        document.add_paragraph(caption)
    return table


def texts(document: Document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


@pytest.fixture()
def profile(profile_file: Path):
    return load_profile_from_path(profile_file)


def test_figure_caption_above_is_moved_below(make_context, profile) -> None:
    document = Document()
    add_figure(document, "Рисунок 1 – Схема", caption_first=True)

    context = make_context(document, profile)
    moves = apply_caption_positions(context)

    assert moves == 1
    assert texts(document) == ["", "Рисунок 1 – Схема"]


def test_figure_caption_below_stays(make_context, profile) -> None:
    document = Document()
    add_figure(document, "Рисунок 1 – Схема")

    context = make_context(document, profile)
    moves = apply_caption_positions(context)

    assert moves == 0
    assert texts(document) == ["", "Рисунок 1 – Схема"]


def test_blank_line_between_does_not_prevent_move(make_context, profile) -> None:
    document = Document()
    document.add_paragraph("Рисунок 1 – Схема")
    document.add_paragraph("")
    add_figure(document)

    context = make_context(document, profile)
    assert apply_caption_positions(context) == 1

    # Пустая строка осталась над рисунком, подпись ушла под него.
    assert texts(document) == ["", "", "Рисунок 1 – Схема"]


def test_table_caption_below_is_moved_above(make_context, profile) -> None:
    document = Document()
    add_table(document, "Таблица 1 – Данные")

    context = make_context(document, profile)
    moves = apply_caption_positions(context)

    assert moves == 1
    # Название должно оказаться до таблицы, то есть первым абзацем документа
    assert texts(document)[0] == "Таблица 1 – Данные"


def test_table_caption_above_stays(make_context, profile) -> None:
    document = Document()
    add_table(document, "Таблица 1 – Данные", caption_first=True)

    context = make_context(document, profile)
    assert apply_caption_positions(context) == 0


def test_move_is_reported(make_context, profile) -> None:
    document = Document()
    add_figure(document, "Рисунок 1 – Схема", caption_first=True)

    context = make_context(document, profile)
    apply_caption_positions(context)

    moves = [
        change
        for change in context.changes.by_category()["caption-positions"]
        if change.description == "подпись рисунка перенесена под рисунок"
    ]
    assert len(moves) == 1
    assert moves[0].text == "Рисунок 1 – Схема"
    assert "было: над рисунком" in moves[0].detail
    assert "перенесено 1" in context.changes.items[-1].description


def test_two_captions_are_not_guessed(make_context, profile) -> None:
    """Подписи и над рисунком, и под ним — приложение не угадывает."""
    document = Document()
    document.add_paragraph("Рисунок 1 – Схема")
    add_figure(document, "Рисунок 2 – Другая схема")

    context = make_context(document, profile)
    moves = apply_caption_positions(context)

    assert moves == 0
    assert texts(document) == ["Рисунок 1 – Схема", "", "Рисунок 2 – Другая схема"]
    manual = context.changes.by_category()["caption-positions-manual"]
    assert manual and "две подписи" in manual[0].description


def test_caption_of_another_figure_is_not_taken(make_context, profile) -> None:
    """Подпись между двумя рисунками принадлежит первому — второй её не забирает."""
    document = Document()
    add_figure(document)
    document.add_paragraph("Рисунок 1 – Схема")
    add_figure(document)

    context = make_context(document, profile)
    moves = apply_caption_positions(context)

    assert moves == 0
    assert texts(document) == ["", "Рисунок 1 – Схема", ""]


def test_position_keep_does_nothing(make_context, profile_dict: dict, tmp_path: Path) -> None:
    data = copy.deepcopy(profile_dict)
    data["captions"]["figure"]["position"] = "keep"
    data["captions"]["table"]["position"] = "keep"
    path = tmp_path / "не-переносить.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    add_figure(document, "Рисунок 1 – Схема", caption_first=True)

    context = make_context(document, load_profile_from_path(path))
    assert apply_caption_positions(context) == 0
    assert texts(document) == ["Рисунок 1 – Схема", ""]


def test_caption_is_moved_above_when_profile_requires(
    make_context, profile_dict: dict, tmp_path: Path
) -> None:
    data = copy.deepcopy(profile_dict)
    data["captions"]["figure"]["position"] = "above"
    path = tmp_path / "подпись-сверху.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    add_figure(document, "Рисунок 1 – Схема")

    context = make_context(document, load_profile_from_path(path))
    assert apply_caption_positions(context) == 1
    assert texts(document) == ["Рисунок 1 – Схема", ""]


def test_text_and_tables_are_not_moved(make_context, profile) -> None:
    """Кроме подписей ничего не двигается."""
    document = Document()
    document.add_paragraph("Первый абзац текста.")
    add_table(document, "Таблица 1 – Данные", caption_first=True)
    document.add_paragraph("Последний абзац текста.")

    context = make_context(document, profile)
    assert apply_caption_positions(context) == 0
    assert texts(document) == [
        "Первый абзац текста.",
        "Таблица 1 – Данные",
        "Последний абзац текста.",
    ]


def test_second_run_moves_nothing(make_context, profile) -> None:
    document = Document()
    add_figure(document, "Рисунок 1 – Схема", caption_first=True)

    first = make_context(document, profile)
    apply_caption_positions(first)
    after_first = texts(document)

    second = make_context(document, profile)
    assert apply_caption_positions(second) == 0
    assert texts(document) == after_first
