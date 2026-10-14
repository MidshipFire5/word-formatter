"""Правила оформления подписей рисунков и названий таблиц."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_captions


def test_figure_caption_is_centered(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Рисунок 1 – Схема процесса")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert paragraph.paragraph_format.line_spacing == 1.0
    run = paragraph.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)


def test_table_caption_has_no_indent(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Таблица 1 – Показатели деятельности")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert paragraph.runs[0].font.size == Pt(14)


def test_continued_table_caption_is_formatted_as_table(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Продолжение таблицы 2.4")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    assert paragraph.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert paragraph.runs[0].font.size == Pt(14)


def test_captions_forbid_hyphenation(make_context, profile_file: Path) -> None:
    document = Document()
    figure = document.add_paragraph("Рисунок 2 – Название")
    table = document.add_paragraph("Таблица 2 – Название")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    for paragraph in (figure, table):
        assert paragraph._p.pPr.find(qn("w:suppressAutoHyphenation")) is not None


def test_body_paragraph_is_not_touched(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Обычный текст документа.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    assert paragraph.paragraph_format.alignment is None
    assert context.changes.count("figure-captions") == 0
    assert context.changes.count("table-captions") == 0


def test_captions_are_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Рисунок 1 – Первый")
    document.add_paragraph("Рисунок 2 – Второй")
    document.add_paragraph("Таблица 1 – Данные")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_captions(context)

    assert "подписей: 2" in context.changes.by_category()["figure-captions"][0].detail
    assert "названий: 1" in context.changes.by_category()["table-captions"][0].detail
