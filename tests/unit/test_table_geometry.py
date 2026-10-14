"""Ширина и отступ таблицы.

Таблица уезжает вправо, если у неё задана ширина (галочка «Ширина»
в свойствах таблицы) или отступ от левого поля. Профиль может снять и то,
и другое.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_tables
from word_formatter.core.rules.table_geometry import set_table_geometry


@pytest.fixture()
def profile(profile_file: Path):
    return load_profile_from_path(profile_file)


def build_table(document: Document, *, width: str | None = '9639', indent: str | None = "709"):
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Выручка"
    table.cell(1, 1).text = "55000"

    tblPr = table._tbl.tblPr
    if width is not None:
        tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{width}" w:type="dxa"/>'))
    if indent is not None:
        tblPr.append(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{indent}" w:type="dxa"/>'))
    return table


def widths_of(table) -> list[tuple[str | None, str | None]]:
    return [
        (node.get(qn("w:w")), node.get(qn("w:type")))
        for node in table._tbl.tblPr.findall(qn("w:tblW"))
    ]


def indent_of(table) -> str | None:
    node = table._tbl.tblPr.find(qn("w:tblInd"))
    return None if node is None else node.get(qn("w:w"))


def test_width_is_removed(make_context, profile) -> None:
    """Снятая «Ширина» — это w:tblW со значением auto."""
    document = Document()
    table = build_table(document)

    apply_tables(make_context(document, profile))

    assert widths_of(table) == [("0", "auto")]


def test_duplicate_width_elements_are_collapsed(make_context, profile) -> None:
    """Word может сохранить два w:tblW — остаться должен один."""
    document = Document()
    table = build_table(document, indent=None)
    table._tbl.tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="5000" w:type="pct"/>'))
    assert len(widths_of(table)) > 1

    apply_tables(make_context(document, profile))

    assert widths_of(table) == [("0", "auto")]


def test_indent_is_removed(make_context, profile) -> None:
    document = Document()
    table = build_table(document)
    assert indent_of(table) == "709"

    apply_tables(make_context(document, profile))

    assert indent_of(table) == "0"


def test_indent_is_not_duplicated(make_context, profile) -> None:
    document = Document()
    table = build_table(document)

    apply_tables(make_context(document, profile))

    assert len(table._tbl.tblPr.findall(qn("w:tblInd"))) == 1


def test_keep_mode_leaves_geometry_alone(make_context, profile_dict: dict, tmp_path: Path) -> None:
    """Профиль без требований к положению таблицы ничего не меняет."""
    data = copy.deepcopy(profile_dict)
    data["tables"]["preferredWidth"] = "keep"
    data["tables"].pop("indentCm")
    path = tmp_path / "не-менять-таблицы.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    table = build_table(document)
    before = widths_of(table)

    apply_tables(make_context(document, load_profile_from_path(path)))

    assert widths_of(table) == before
    assert indent_of(table) == "709"


def test_table_without_properties_gets_them(make_context, profile) -> None:
    """Свойства таблицы могут отсутствовать — тогда их нужно создать."""
    document = Document()
    table = build_table(document)
    tbl = table._tbl
    tbl.remove(tbl.find(qn("w:tblPr")))
    assert tbl.find(qn("w:tblPr")) is None

    apply_tables(make_context(document, profile))

    tblPr = tbl.find(qn("w:tblPr"))
    assert tblPr is not None
    assert tbl[0] is tblPr, "w:tblPr обязан идти первым в w:tbl"
    assert widths_of(table) == [("0", "auto")]


def test_properties_order_matches_schema(make_context, profile) -> None:
    """Порядок внутри w:tblPr задан схемой: сначала ширина, потом отступ."""
    document = Document()
    table = build_table(document, width=None, indent=None)

    apply_tables(make_context(document, profile))

    names = [child.tag.rsplit("}", 1)[-1] for child in table._tbl.tblPr]
    assert names.index("tblW") < names.index("tblInd")


def test_custom_indent_is_applied(profile_dict: dict) -> None:
    """Профиль может задать отступ числом — тогда он и выставляется."""
    data = copy.deepcopy(profile_dict)
    data["tables"]["indentCm"] = 1.5
    settings = load_profile_from_path_dict(data).tables

    document = Document()
    table = build_table(document)

    actions = set_table_geometry(table, settings)

    assert indent_of(table) == str(Cm(1.5).twips)
    assert any("отступ" in action for action in actions)


def load_profile_from_path_dict(data: dict):
    """Собирает профиль из словаря через настоящую загрузку."""
    import json as json_module
    import tempfile

    with tempfile.NamedTemporaryFile(
        "w", suffix=".json", delete=False, encoding="utf-8"
    ) as handle:
        json_module.dump(data, handle, ensure_ascii=False)
        name = handle.name
    return load_profile_from_path(Path(name))


def test_position_is_reported(make_context, profile) -> None:
    document = Document()
    build_table(document)

    context = make_context(document, profile)
    apply_tables(context)

    reports = [
        change
        for change in context.changes.by_category()["tables"]
        if change.description.startswith("положение таблиц")
    ]
    assert len(reports) == 1
    assert "снята заданная ширина" in reports[0].description
    assert "убран отступ от левого поля" in reports[0].description
