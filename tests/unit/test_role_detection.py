"""Определение роли абзаца: включая случаи, где приложение обязано не угадывать."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

from word_formatter.core.docx import Role, classify_paragraph
from word_formatter.core.profiles import load_profile_from_path


@pytest.fixture()
def profile(profile_file: Path):
    return load_profile_from_path(profile_file)


def build_document(tmp_path: Path, specs: list[dict]) -> Document:
    document = Document()
    for spec in specs:
        paragraph = document.add_paragraph(spec.get("text", ""), style=spec.get("style"))
        if spec.get("bold"):
            for run in paragraph.runs:
                run.bold = True
        if spec.get("center"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if spec.get("italic"):
            for run in paragraph.runs:
                run.italic = True
        if spec.get("raw"):
            paragraph._p.append(parse_xml(spec["raw"]))
    return document


def role_of(tmp_path: Path, profile, spec: dict) -> tuple[Role, int | None, str]:
    document = build_document(tmp_path, [spec])
    info = classify_paragraph(document.paragraphs[0], profile)
    return info.role, info.level, info.reason


@pytest.mark.parametrize(
    ("text", "expected_role"),
    [
        ("ВВЕДЕНИЕ", Role.STRUCTURAL_HEADING),
        ("Введение", Role.STRUCTURAL_HEADING),
        ("ЗАКЛЮЧЕНИЕ", Role.STRUCTURAL_HEADING),
        ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", Role.STRUCTURAL_HEADING),
        ("СОДЕРЖАНИЕ", Role.STRUCTURAL_HEADING),
        ("Приложение А", Role.STRUCTURAL_HEADING),
        ("ПРИЛОЖЕНИЕ 1", Role.STRUCTURAL_HEADING),
        ("Рисунок 1 – Схема процесса", Role.FIGURE_CAPTION),
        ("Рисунок 2.1 — Название", Role.FIGURE_CAPTION),
        ("Рис. 3 – Схема", Role.FIGURE_CAPTION),
        ("Таблица 1 – Показатели деятельности", Role.TABLE_CAPTION),
        ("Таблица 2.4 – Расходы, руб.", Role.TABLE_CAPTION),
        ("Продолжение таблицы 2.4", Role.TABLE_CONTINUATION),
        ("1 Теоретические основы процесса автоматизации", Role.HEADING),
        ("1.1 Понятие и сущность процесса", Role.HEADING),
        ("1.1.2 Особенности подхода", Role.HEADING),
        ("2020 год был трудным для отрасли", Role.BODY),
        ("1,5 интервала в тексте", Role.BODY),
        ("5 таблиц приведено в работе", Role.BODY),
        ("Текст, текст, текст.", Role.BODY),
        ("", Role.BLANK),
        ("   ", Role.BLANK),
        ("- первый пункт перечисления", Role.LIST_ITEM),
        ("– второй пункт перечисления", Role.LIST_ITEM),
        ("• третий пункт перечисления", Role.LIST_ITEM),
    ],
)
def test_role_is_detected(tmp_path: Path, profile, text: str, expected_role: Role) -> None:
    role, _level, _reason = role_of(tmp_path, profile, {"text": text})
    assert role is expected_role


@pytest.mark.parametrize(
    ("text", "expected_level"),
    [
        ("1 Теоретические основы", 1),
        ("1.1 Понятие и сущность", 2),
        ("1.1.2 Особенности подхода", 3),
        ("2.10 Последний подраздел главы", 2),
    ],
)
def test_heading_level_comes_from_number(tmp_path: Path, profile, text: str, expected_level: int) -> None:
    _role, level, _reason = role_of(tmp_path, profile, {"text": text})
    assert level == expected_level


def test_heading_style_gives_level(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "Название раздела", "style": "Heading 1"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.HEADING
    assert info.level == 1
    assert info.style_name


def test_second_level_style_gives_level_two(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "Подраздел", "style": "Heading 2"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.HEADING
    assert info.level == 2


def test_structural_heading_wins_over_style(tmp_path: Path, profile) -> None:
    """«ЗАКЛЮЧЕНИЕ» остаётся структурным заголовком, даже оформленное стилем раздела."""
    document = build_document(tmp_path, [{"text": "ЗАКЛЮЧЕНИЕ", "style": "Heading 1"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.STRUCTURAL_HEADING


def test_heading_style_with_body_text_is_ambiguous(tmp_path: Path, profile) -> None:
    """Длинный текст с точкой в конце, но со стилем заголовка — не оформляем и сообщаем."""
    long_text = "Этот абзац оформлен стилем заголовка, но по содержанию это обычный текст. " * 3
    document = build_document(tmp_path, [{"text": long_text.strip(), "style": "Heading 1"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.AMBIGUOUS
    assert "похож на обычный текст" in info.reason


def test_numbered_bold_paragraph_with_sentence_is_ambiguous(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "1 Показатель равен 5.", "bold": True}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.AMBIGUOUS
    assert "признаки обычного текста" in info.reason


def test_list_style_is_list_item(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "Элемент списка", "style": "List Bullet"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.LIST_ITEM


def test_numbered_list_is_list_item_not_heading(tmp_path: Path, profile) -> None:
    """Автонумерованный список не должен приниматься за раздел."""
    document = build_document(tmp_path, [{"text": "Первый пункт", "style": "List Number"}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.LIST_ITEM


def test_formula_paragraph_is_detected(tmp_path: Path, profile) -> None:
    raw = (
        f'<m:oMathPara {nsdecls("m")}><m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></m:oMathPara>'
    )
    document = build_document(tmp_path, [{"text": "", "raw": raw}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.has_formula is True
    assert info.role is Role.FORMULA


def test_embedded_object_is_not_touched(tmp_path: Path, profile) -> None:
    raw = (
        '<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape id="_x0000_i1025" type="#_x0000_t75"/></w:object>'
    )
    document = build_document(tmp_path, [{"text": "", "raw": raw}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.AMBIGUOUS
    assert "встроенный объект" in info.reason


@pytest.mark.parametrize(
    "text",
    [
        "1 Теоретические основы ……………………………………… 4",
        "1.1 Понятие и сущность ....................... 5",
        "ВВЕДЕНИЕ …………………………………………………………… 3",
        "ЗАКЛЮЧЕНИЕ . . . . . . . . . . . . . . . . . 20",
    ],
)
def test_contents_line_is_not_a_heading(tmp_path: Path, profile, text: str) -> None:
    """Строка содержания с отточием — не заголовок, даже если начинается с номера."""
    role, _level, _reason = role_of(tmp_path, profile, {"text": text})
    assert role is Role.BODY


def test_paragraph_inside_table_has_cell_role(tmp_path: Path, profile) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "55000"
    info = classify_paragraph(table.cell(0, 0).paragraphs[0], profile, in_table=True, table_index=0)
    assert info.role is Role.TABLE_CELL
    assert info.in_table is True
    assert info.table_index == 0


def test_italic_body_text_stays_body(tmp_path: Path, profile) -> None:
    """Курсив — оформление, а не роль: абзац остаётся основным текстом."""
    document = build_document(tmp_path, [{"text": "Текст с курсивом внутри.", "italic": True}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.BODY


def test_centered_short_title_is_not_heading(tmp_path: Path, profile) -> None:
    """Титульная строка по центру не становится заголовком без стиля и номера."""
    document = build_document(
        tmp_path, [{"text": "КРАСНОДАРСКИЙ ФИЛИАЛ ФИНУНИВЕРСИТЕТА", "center": True, "bold": True}]
    )
    info = classify_paragraph(document.paragraphs[0], profile)
    assert info.role is Role.BODY


def test_short_text_is_prepared_for_report(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "Слово " * 40}])
    info = classify_paragraph(document.paragraphs[0], profile)
    assert len(info.short_text) <= 80
    assert info.short_text.endswith("…")


def test_paragraph_index_and_style_are_recorded(tmp_path: Path, profile) -> None:
    document = build_document(tmp_path, [{"text": "Обычный текст.", "style": "Normal"}])
    info = classify_paragraph(document.paragraphs[0], profile, index=7)
    assert info.index == 7
    assert info.style_name == "Normal"
    assert info.has_drawing is False
    assert info.has_formula is False
