"""Проверка сохранности содержимого."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from word_formatter.core.docx import open_document
from word_formatter.core.verification import (
    VerificationResult,
    verify_documents,
    verify_files,
)


def test_identical_documents_pass() -> None:
    source = Document()
    source.add_paragraph("Текст документа.")
    result = Document()
    result.add_paragraph("Текст документа.")

    check = verify_documents(source, result)

    assert isinstance(check, VerificationResult)
    assert check.passed is True
    assert check.problems == ()


def test_changed_text_is_reported() -> None:
    source = Document()
    source.add_paragraph("Исходный текст.")
    result = Document()
    result.add_paragraph("Изменённый текст.")

    check = verify_documents(source, result)

    assert check.passed is False
    assert any("текст" in problem for problem in check.problems)


def test_added_empty_paragraph_is_not_a_loss() -> None:
    source = Document()
    source.add_paragraph("Текст.")
    result = Document()
    result.add_paragraph("Текст.")
    result.add_paragraph("")

    check = verify_documents(source, result)

    assert check.passed is True
    assert check.notes and "пустых абзацев" in check.notes[0]


def test_changed_table_text_is_reported() -> None:
    source = Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "55000"

    result = Document()
    table = result.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "54000"

    check = verify_documents(source, result)

    assert check.passed is False
    assert any("таблиц" in problem for problem in check.problems)


def test_order_change_without_moves_is_a_difference() -> None:
    """Порядок абзацев — часть содержания: без причины он меняться не должен."""
    source = Document()
    source.add_paragraph("Первый абзац.")
    source.add_paragraph("Второй абзац.")
    result = Document()
    result.add_paragraph("Второй абзац.")
    result.add_paragraph("Первый абзац.")

    check = verify_documents(source, result)

    assert check.passed is False
    assert any("порядок" in problem for problem in check.problems)


def test_order_change_with_expected_moves_is_a_note() -> None:
    """Перенос подписи меняет порядок двух абзацев — об этом сообщается, но это не ошибка."""
    source = Document()
    source.add_paragraph("Рисунок 1 – Схема")
    source.add_paragraph("Текст после рисунка.")
    result = Document()
    result.add_paragraph("Текст после рисунка.")
    result.add_paragraph("Рисунок 1 – Схема")

    check = verify_documents(source, result, expected_moves=1)

    assert check.passed is True, check.problems
    assert any("порядок абзацев изменён" in note for note in check.notes)


def test_lost_text_is_reported_even_with_moves() -> None:
    source = Document()
    source.add_paragraph("Первый абзац.")
    source.add_paragraph("Второй абзац.")
    result = Document()
    result.add_paragraph("Второй абзац.")

    check = verify_documents(source, result, expected_moves=1)

    assert check.passed is False
    assert any("текст" in problem for problem in check.problems)


def test_missing_image_is_reported(rich_docx: Path) -> None:
    source = open_document(rich_docx)
    result = open_document(rich_docx)

    drawing = result.element.body.find(".//" + qn("w:drawing"))
    assert drawing is not None
    drawing.getparent().remove(drawing)

    check = verify_documents(source, result)

    assert check.passed is False
    assert any("изображений" in problem for problem in check.problems)


def test_missing_footnote_is_reported(rich_docx: Path) -> None:
    source = open_document(rich_docx)
    result = open_document(rich_docx)

    reference = result.element.body.find(".//" + qn("w:footnoteReference"))
    assert reference is not None
    reference.getparent().remove(reference)

    check = verify_documents(source, result)

    assert check.passed is False
    assert any("сносок" in problem for problem in check.problems)


def test_rich_document_matches_itself(rich_docx: Path) -> None:
    check = verify_files(rich_docx, rich_docx)
    assert check.passed is True
