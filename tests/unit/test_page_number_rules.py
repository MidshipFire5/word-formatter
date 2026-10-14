"""Правила нумерации страниц."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from word_formatter.core.docx import open_document
from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_page_numbers
from word_formatter.core.rules.page_numbers import has_page_field, split_position


def test_position_is_split() -> None:
    assert split_position("bottom-center") == ("bottom", "center")
    assert split_position("top-right") == ("top", "right")


def test_field_is_added_to_footer(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    footer = document.sections[0].footer
    assert footer.is_linked_to_previous is False
    paragraph = footer.paragraphs[0]
    assert has_page_field(paragraph)
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.runs[0].font.size == Pt(12)
    assert paragraph.runs[0].font.name == "Times New Roman"


def test_field_instruction_is_page(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    instructions = [
        node.text
        for node in document.sections[0].footer.paragraphs[0]._p.iter(qn("w:instrText"))
    ]
    assert any("PAGE" in (text or "") for text in instructions)


def test_existing_field_is_updated_not_duplicated(rich_docx: Path, profile_file: Path) -> None:
    from word_formatter.core.docx import inspect_document
    from word_formatter.core.rules import RuleContext

    document = open_document(rich_docx)
    profile = load_profile_from_path(profile_file)
    context = RuleContext(
        document=document, profile=profile, inventory=inspect_document(document, profile)
    )
    before = len(document.sections[0].footer.paragraphs)

    apply_page_numbers(context)

    footer = document.sections[0].footer
    assert len(footer.paragraphs) == before
    fields = [
        node
        for paragraph in footer.paragraphs
        for node in paragraph._p.iter(qn("w:instrText"))
        if "PAGE" in (node.text or "").upper()
    ]
    assert len(fields) == 1
    assert "оформление номера обновлено: 1" in context.changes.by_category()["page-numbers"][0].detail


def test_first_page_has_no_number(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    assert document.sections[0].different_first_page_header_footer is True
    # Поле стоит в обычном колонтитуле, а не в колонтитуле первой страницы
    assert has_page_field(document.sections[0].footer.paragraphs[0])


def test_second_section_inherits_footer(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_section()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    assert len(document.sections) == 2
    assert document.sections[1].footer.is_linked_to_previous is True
    fields = [
        node
        for paragraph in document.sections[1].footer.paragraphs
        for node in paragraph._p.iter(qn("w:instrText"))
        if "PAGE" in (node.text or "").upper()
    ]
    assert len(fields) == 1, "во втором разделе не должно появиться второго номера"


def test_top_position_uses_header(make_context, profile_dict: dict, tmp_path: Path) -> None:
    data = copy.deepcopy(profile_dict)
    data["pageNumbers"]["position"] = "top-right"
    path = tmp_path / "сверху.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    context = make_context(document, load_profile_from_path(path))
    apply_page_numbers(context)

    assert has_page_field(document.sections[0].header.paragraphs[0])
    assert document.sections[0].header.paragraphs[0].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_numbering_restart_is_removed(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_section()
    sectPr = document.sections[1]._sectPr
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    sectPr.append(pgNumType)

    context = make_context(document, load_profile_from_path(profile_file))
    apply_page_numbers(context)

    pgNumType = document.sections[1]._sectPr.find(qn("w:pgNumType"))
    assert pgNumType.get(qn("w:start")) is None
    assert pgNumType.get(qn("w:fmt")) == "decimal"
    assert "снято сбросов нумерации: 1" in context.changes.by_category()["page-numbers"][0].detail


def test_fields_are_marked_for_update(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    node = document.settings.element.find(qn("w:updateFields"))
    assert node is not None
    assert node.get(qn("w:val")) == "true"


def test_disabled_profile_is_reported(make_context, profile_dict: dict, tmp_path: Path) -> None:
    data = copy.deepcopy(profile_dict)
    data["pageNumbers"]["enabled"] = False
    path = tmp_path / "без-номеров.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    context = make_context(document, load_profile_from_path(path))
    apply_page_numbers(context)

    assert "отключена" in context.changes.by_category()["page-numbers"][0].description


def test_report_mentions_title_page_rule(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_page_numbers(context)

    description = context.changes.by_category()["page-numbers"][0].description
    assert "внизу по центру" in description
    assert "титульном листе" in description
