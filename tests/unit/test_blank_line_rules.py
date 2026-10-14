"""Правила вставки свободных строк."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm, Pt

from word_formatter.core.docx import Role
from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_blank_lines
from word_formatter.core.rules.blank_lines import (
    apply_blank_lines,
    belongs_together,
    needs_blank_line,
)
from word_formatter.core.rules.body import BodyElement


def item(role: Role, kind: str = "paragraph") -> BodyElement:
    return BodyElement(kind=kind, index=0, role=role, element=None, text="")


TABLE = item(Role.BODY, kind="table")


@pytest.mark.parametrize(
    ("previous", "following", "expected"),
    [
        (item(Role.STRUCTURAL_HEADING), item(Role.BODY), True),
        (item(Role.HEADING), item(Role.BODY), True),
        (item(Role.BODY), item(Role.HEADING), True),
        (item(Role.HEADING), item(Role.STRUCTURAL_HEADING), True),
        (item(Role.BODY), item(Role.BODY), False),
        (item(Role.LIST_ITEM), item(Role.BODY), False),
        (item(Role.BODY), item(Role.FORMULA), True),
        (item(Role.FORMULA), item(Role.BODY), True),
        (item(Role.FIGURE), item(Role.FIGURE_CAPTION), False),
        (item(Role.FIGURE_CAPTION), item(Role.BODY), True),
        (item(Role.BODY), item(Role.FIGURE), True),
        (item(Role.TABLE_CAPTION), TABLE, False),
        (TABLE, item(Role.BODY), True),
        (item(Role.BODY), TABLE, True),
        (TABLE, TABLE, False),
        (TABLE, item(Role.TABLE_CONTINUATION), False),
        (item(Role.BLANK), item(Role.HEADING), False),
        (item(Role.HEADING), item(Role.BLANK), False),
        (item(Role.AMBIGUOUS), item(Role.BODY), False),
        (item(Role.BODY), item(Role.AMBIGUOUS), False),
    ],
)
def test_blank_line_decision(previous, following, expected: bool) -> None:
    assert needs_blank_line(previous, following) is expected


def test_parts_of_one_block_belong_together() -> None:
    assert belongs_together(item(Role.FIGURE), item(Role.FIGURE_CAPTION)) is True
    assert belongs_together(item(Role.TABLE_CAPTION), TABLE) is True
    assert belongs_together(TABLE, TABLE) is True
    assert belongs_together(item(Role.BODY), item(Role.BODY)) is False


def texts(document: Document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def test_blank_line_between_heading_and_text(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("Текст введения.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["ВВЕДЕНИЕ", "", "Текст введения."]
    assert context.changes.count("blank-lines") == 2  # вставка + итог


def test_blank_line_between_text_and_heading(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Текст раздела.")
    document.add_paragraph("1.1 Новый подраздел")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["Текст раздела.", "", "1.1 Новый подраздел"]


def test_blank_line_between_two_headings(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("1 Раздел")
    document.add_paragraph("1.1 Подраздел")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["1 Раздел", "", "1.1 Подраздел"]


def test_blank_lines_around_table(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Текст до таблицы.")
    document.add_paragraph("Таблица 1 – Данные")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(1, 0).text = "Выручка"
    document.add_paragraph("Текст после таблицы.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == [
        "Текст до таблицы.",
        "",
        "Таблица 1 – Данные",
        "",
        "Текст после таблицы.",
    ]


def test_blank_line_is_not_inserted_twice(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("")
    document.add_paragraph("Текст введения.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["ВВЕДЕНИЕ", "", "Текст введения."]
    assert "вставлено 0" in context.changes.items[-1].description


def test_repeated_run_adds_nothing(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("Текст введения.")
    document.add_paragraph("1.1 Подраздел")
    document.add_paragraph("Текст подраздела.")
    profile = load_profile_from_path(profile_file)

    first = make_context(document, profile)
    apply_blank_lines(first)
    after_first = texts(document)

    second = make_context(document, profile)
    apply_blank_lines(second)

    assert texts(document) == after_first
    assert "вставлено 0" in second.changes.items[-1].description


def test_no_blank_line_at_document_start(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["ВВЕДЕНИЕ"]


def test_no_blank_line_before_page_break(make_context, profile_file: Path) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph("Текст предыдущего раздела.")
    paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:br w:type="page"/></w:r>'))
    document.add_paragraph("ВВЕДЕНИЕ")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["Текст предыдущего раздела.", "ВВЕДЕНИЕ"]


def test_no_blank_line_before_paragraph_starting_new_page(
    make_context, profile_file: Path
) -> None:
    document = Document()
    document.add_paragraph("Текст предыдущего раздела.")
    heading = document.add_paragraph("ВВЕДЕНИЕ")
    heading.paragraph_format.page_break_before = True
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    assert texts(document) == ["Текст предыдущего раздела.", "ВВЕДЕНИЕ"]


def test_inserted_paragraph_is_formatted_as_body(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("Текст введения.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    blank = document.paragraphs[1]
    assert blank.text == ""
    assert blank.paragraph_format.line_spacing == 1.5
    assert blank.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert blank.paragraph_format.space_after == Pt(0)


def test_insertions_are_reported_with_context(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("1 Раздел")
    document.add_paragraph("Текст раздела.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    insertions = [
        change
        for change in context.changes.by_category()["blank-lines"]
        if change.description == "вставлена свободная строка"
    ]
    assert len(insertions) == 1
    assert insertions[0].text == "Текст раздела."
    assert "заголовок раздела" in insertions[0].detail


def test_disabled_profile_is_reported(make_context, profile_dict: dict, tmp_path: Path) -> None:
    data = copy.deepcopy(profile_dict)
    data["blankLines"]["insert"] = False
    path = tmp_path / "без-строк.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    document.add_paragraph("ВВЕДЕНИЕ")
    document.add_paragraph("Текст введения.")
    context = make_context(document, load_profile_from_path(path))
    apply_blank_lines(context)

    assert texts(document) == ["ВВЕДЕНИЕ", "Текст введения."]
    assert "отключена" in context.changes.by_category()["blank-lines"][0].description


def test_ambiguous_paragraph_does_not_trigger_insertion(make_context, profile_file: Path) -> None:
    """Спорные абзацы не оформляются и не обрастают пустыми строками."""
    from docx.oxml import parse_xml

    document = Document()
    document.add_paragraph("Текст документа.")
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            '<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml">'
            '<v:shape id="_x0000_i1025"/></w:object>'
        )
    )
    document.add_paragraph("Продолжение текста.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_blank_lines(context)

    insertions = [
        change
        for change in context.changes.by_category()["blank-lines"]
        if change.description == "вставлена свободная строка"
    ]
    assert insertions == []
