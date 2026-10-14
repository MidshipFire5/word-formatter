"""Правила оформления перечислений."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_lists


def test_manual_marker_item_is_formatted(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("– первый пункт перечисления")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_lists(context)

    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(1.25).twips
    run = paragraph.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)
    # Маркер остаётся частью текста: приложение его не переписывает
    assert paragraph.text == "– первый пункт перечисления"


def test_styled_list_item_is_formatted_and_keeps_style(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Элемент списка", style="List Bullet")
    paragraph.runs[0].font.size = Pt(9)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_lists(context)

    assert paragraph.runs[0].font.size == Pt(14)
    assert paragraph.style.name == "List Bullet"


def test_emphasis_in_list_item_is_removed(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("– пункт с курсивом")
    paragraph.runs[0].italic = True
    context = make_context(document, load_profile_from_path(profile_file))

    apply_lists(context)

    assert paragraph.runs[0].italic is False
    assert context.changes.count("italic-removed") == 1


def test_body_paragraph_is_not_touched(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Обычный текст без маркера.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_lists(context)

    assert paragraph.paragraph_format.alignment is None
    assert context.changes.count("lists") == 0


def test_markers_are_not_replaced_and_it_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("– пункт первый")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_lists(context)

    assert context.changes.count("list-markers") == 1
    message = context.changes.by_category()["list-markers"][0].description
    assert "не изменялись" in message

    assert "элементов: 1" in context.changes.by_category()["lists"][0].detail
