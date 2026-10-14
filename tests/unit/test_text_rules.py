"""Правила оформления основного текста и оформление по умолчанию."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_body_text, apply_document_defaults

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def make_body_document() -> Document:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Текст документа с курсивом и подчёркиванием.")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.italic = True
    run.underline = True
    return document


def test_body_text_is_reformatted(make_context, profile_file: Path) -> None:
    document = make_body_document()
    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    paragraph = document.paragraphs[0]
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    # Размеры в DOCX хранятся в твипах, поэтому сравниваем в них же.
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(1.25).twips
    assert paragraph.paragraph_format.line_spacing == 1.5
    assert paragraph.paragraph_format.space_before == Pt(0)
    assert paragraph.paragraph_format.space_after == Pt(0)

    run = paragraph.runs[0]
    assert run.font.name == "Times New Roman"
    assert run.font.size == Pt(14)
    assert run.italic is False
    assert run.underline is False


def test_theme_font_reference_is_removed(make_context, profile_file: Path) -> None:
    """Ссылка на шрифт темы перебивает явное имя шрифта, поэтому её снимают."""
    document = make_body_document()
    run = document.paragraphs[0].runs[0]
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    rFonts.set(qn("w:asciiTheme"), "minorHAnsi")
    rFonts.set(qn("w:hAnsiTheme"), "minorHAnsi")

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    rFonts = run._element.rPr.find(qn("w:rFonts"))
    assert rFonts.get(qn("w:asciiTheme")) is None
    assert rFonts.get(qn("w:hAnsiTheme")) is None
    assert rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert rFonts.get(qn("w:hAnsi")) == "Times New Roman"


def test_italic_and_underline_removal_is_reported(make_context, profile_file: Path) -> None:
    document = make_body_document()
    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    assert context.changes.count("italic-removed") == 1
    assert context.changes.count("underline-removed") == 1
    italic_change = context.changes.by_category()["italic-removed"][0]
    assert italic_change.paragraph_index == 0
    assert "Текст документа" in italic_change.text


def test_plain_text_without_emphasis_is_not_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Обычный текст без выделений.")
    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    assert context.changes.count("italic-removed") == 0
    assert context.changes.count("underline-removed") == 0


def test_heading_and_list_keep_their_own_formatting(make_context, profile_file: Path) -> None:
    """На этом шаге оформляются только абзацы основного текста."""
    document = Document()
    heading = document.add_paragraph("Название раздела", style="Heading 1")
    heading_run = heading.add_run("")
    item = document.add_paragraph("Элемент списка", style="List Bullet")

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    assert heading.paragraph_format.alignment is None
    assert item.paragraph_format.alignment is None
    assert heading_run.font.size is None


def test_footnote_reference_run_is_untouched(make_context, profile_file: Path) -> None:
    """Прогон со сноской не содержит текста: менять ему шрифт нельзя."""
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Текст со сноской")
    reference = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
        '<w:footnoteReference w:id="1"/></w:r>'
    )
    paragraph._p.append(reference)

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    rPr = reference.find(qn("w:rPr"))
    assert rPr.find(qn("w:rFonts")) is None
    assert rPr.find(qn("w:sz")) is None
    assert rPr.find(qn("w:rStyle")) is not None


def test_hanging_indent_is_replaced_by_first_line_indent(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Абзац с висячим отступом.")
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    indent = document.paragraphs[0]._p.pPr.find(qn("w:ind"))
    assert indent.get(qn("w:hanging")) is None
    assert indent.get(qn("w:firstLine")) == str(Cm(1.25).twips)


def test_character_based_indent_is_removed(make_context, profile_file: Path) -> None:
    """Отступ в символах перебивает отступ в сантиметрах — его нужно снять."""
    document = Document()
    paragraph = document.add_paragraph("Абзац с отступом в символах.")
    pPr = paragraph._p.get_or_add_pPr()
    indent = parse_xml(f'<w:ind {nsdecls("w")} w:firstLineChars="200" w:firstLine="0"/>')
    pPr.append(indent)

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    indent = document.paragraphs[0]._p.pPr.find(qn("w:ind"))
    assert indent.get(qn("w:firstLineChars")) is None
    assert indent.get(qn("w:firstLine")) == str(Cm(1.25).twips)


def test_automatic_spacing_is_removed(make_context, profile_file: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("Абзац с автоматическими интервалами.")
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:beforeAutospacing="1" w:afterAutospacing="1"/>'))

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    spacing = document.paragraphs[0]._p.pPr.find(qn("w:spacing"))
    assert spacing.get(qn("w:beforeAutospacing")) is None
    assert spacing.get(qn("w:afterAutospacing")) is None
    assert spacing.get(qn("w:after")) == "0"


def test_document_defaults_and_normal_style_are_set(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))
    apply_document_defaults(context)

    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(14)
    assert normal.paragraph_format.line_spacing == 1.5
    assert normal.paragraph_format.space_after == Pt(0)

    settings = document.styles.element
    doc_defaults = settings.find(qn("w:docDefaults"))
    assert doc_defaults is not None
    assert settings.index(doc_defaults) == 0, "docDefaults обязан стоять первым элементом"

    rFonts = doc_defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert rFonts.get(qn("w:asciiTheme")) is None

    sz = doc_defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:sz"))
    assert sz.get(qn("w:val")) == "28"

    spacing = doc_defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing.get(qn("w:line")) == "360"
    assert spacing.get(qn("w:lineRule")) == "auto"
    assert spacing.get(qn("w:after")) == "0"

    assert context.changes.count("document-defaults") == 1


def test_element_order_in_rpr_follows_schema(make_context, profile_file: Path) -> None:
    """Порядок элементов внутри w:rPr задан схемой OOXML, иначе Word ругается."""
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Текст")
    rPr = run._element.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="FF0000"/>'))

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    order = [child.tag.rsplit("}", 1)[-1] for child in run._element.rPr]
    # Схема CT_RPr: rFonts, i, color, sz, szCs, u
    assert order.index("rFonts") < order.index("i")
    assert order.index("i") < order.index("color")
    assert order.index("color") < order.index("sz")
    assert order.index("sz") < order.index("szCs")
    assert order.index("szCs") < order.index("u")


def test_title_page_is_left_untouched(make_context, profile_file: Path) -> None:
    """До первой структурной части — титульный лист: его оформление не задано.

    Методичка описывает основной текст, заголовки и таблицы, но не титульный
    лист. Переформатировать его значит испортить готовую вёрстку автора.
    """
    document = Document()
    title = document.add_paragraph("КРАСНОДАРСКИЙ ФИЛИАЛ ФИНУНИВЕРСИТЕТА")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Курсовая работа")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("ВВЕДЕНИЕ")
    body = document.add_paragraph("Текст введения.")

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert subtitle.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.paragraph_format.line_spacing is None
    assert title.runs[0].font.size is None

    assert body.paragraph_format.line_spacing == 1.5
    assert body.runs[0].font.size == Pt(14)
    assert context.changes.count("title-page") == 1
    assert "титульный лист оставлен без изменений" in context.changes.by_category()["title-page"][0].description


def test_without_structural_heading_nothing_is_skipped(make_context, profile_file: Path) -> None:
    document = Document()
    first = document.add_paragraph("Первый абзац без структурных частей.")
    second = document.add_paragraph("Второй абзац.")

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    for paragraph in (first, second):
        assert paragraph.paragraph_format.line_spacing == 1.5
    assert context.changes.count("title-page") == 0


def test_all_paragraphs_of_body_are_processed(make_context, profile_file: Path) -> None:
    document = Document()
    for number in range(5):
        document.add_paragraph(f"Абзац номер {number}.")

    context = make_context(document, load_profile_from_path(profile_file))
    apply_body_text(context)

    body_change = context.changes.by_category()["body-text"][0]
    assert "обработано абзацев: 5" in body_change.detail
    assert "Times New Roman" in body_change.description
    assert "по ширине" in body_change.description
