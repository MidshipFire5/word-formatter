"""Правила оформления таблиц."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_tables
from word_formatter.core.rules.tables import is_mixed_cell, is_numeric_cell


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("55000", True),
        ("1", True),
        ("1 234,5", True),
        ("1\u00a0234,5", True),
        ("-12,5", True),
        ("12 %", True),
        ("55000 руб.", True),
        ("7 шт.", True),
        ("от 5", True),
        ("0,5", True),
        ("2020-2022", False),
        ("III квартал", False),
        ("нет данных", False),
        ("–", False),
        ("х", False),
        ("", False),
        ("руб.", False),
    ],
)
def test_numeric_cell_detection(text: str, expected: bool) -> None:
    assert is_numeric_cell(text) is expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("5 таблиц", True),
        ("Таблица 1 – Данные", True),
        ("55000", False),
        ("Прочие данные", False),
        ("руб.", False),
    ],
)
def test_mixed_cell_detection(text: str, expected: bool) -> None:
    assert is_mixed_cell(text) is expected


def build_table(document: Document):
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    data = [
        ["Показатель", "Значение", "Единица"],
        ["Выручка", "55000", "руб."],
        ["Прочие данные", "–", "х"],
    ]
    for row_index, row in enumerate(data):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    return table


def alignment_of(table, row: int, column: int):
    return table.cell(row, column).paragraphs[0].paragraph_format.alignment


def test_header_row_is_centered(make_context, profile_file: Path) -> None:
    document = Document()
    table = build_table(document)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    for column in range(3):
        assert alignment_of(table, 0, column) == WD_ALIGN_PARAGRAPH.CENTER


def test_numeric_and_text_cells_are_aligned_by_content(make_context, profile_file: Path) -> None:
    document = Document()
    table = build_table(document)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    assert alignment_of(table, 1, 0) == WD_ALIGN_PARAGRAPH.LEFT  # Выручка
    assert alignment_of(table, 1, 1) == WD_ALIGN_PARAGRAPH.CENTER  # 55000
    assert alignment_of(table, 1, 2) == WD_ALIGN_PARAGRAPH.LEFT  # руб.


def test_marker_cells_keep_their_alignment(make_context, profile_file: Path) -> None:
    document = Document()
    table = build_table(document)
    table.cell(2, 1).paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    assert alignment_of(table, 2, 1) == WD_ALIGN_PARAGRAPH.RIGHT
    assert alignment_of(table, 2, 2) is None


def test_cell_text_formatting(make_context, profile_file: Path) -> None:
    document = Document()
    table = build_table(document)
    run = table.cell(1, 0).paragraphs[0].runs[0]
    run.font.size = Pt(18)
    run.italic = True
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    paragraph = table.cell(1, 0).paragraphs[0]
    assert paragraph.runs[0].font.name == "Times New Roman"
    assert paragraph.runs[0].font.size == Pt(12)
    assert paragraph.runs[0].italic is False
    assert paragraph.paragraph_format.first_line_indent.twips == Cm(0).twips
    assert paragraph.paragraph_format.line_spacing == 1.0


def test_emphasis_removal_in_tables_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    table = build_table(document)
    table.cell(1, 0).paragraphs[0].runs[0].italic = True
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    assert context.changes.count("italic-removed") == 1


def test_table_change_reports_counts(make_context, profile_file: Path) -> None:
    document = Document()
    build_table(document)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    detail = context.changes.by_category()["tables"][0].detail
    assert "таблиц: 1" in detail
    assert "шапка: 3" in detail
    assert "числовых: 1" in detail
    assert "ячеек-прочерков без изменений: 2" in detail


def test_document_without_tables_is_not_reported(make_context, profile_file: Path) -> None:
    document = Document()
    document.add_paragraph("Текст без таблиц.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    assert context.changes.count("tables") == 0


def test_mixed_cell_is_reported_for_manual_check(make_context, profile_file: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Показатель"
    table.cell(1, 0).text = "5 таблиц"
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    messages = [change.description for change in context.changes.by_category()["tables-manual"]]
    assert any("и текст, и числа" in message for message in messages)


def test_merged_cells_are_reported(make_context, profile_file: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Объединённая шапка"
    context = make_context(document, load_profile_from_path(profile_file))

    apply_tables(context)

    messages = [change.description for change in context.changes.by_category()["tables-manual"]]
    assert any("объединённые ячейки" in message for message in messages)
