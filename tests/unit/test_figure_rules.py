"""Оформление рисунков: по центру и без абзацного отступа.

Проверяется главная причина сдвига рисунка вправо: абзац с картинкой обычно
не имеет собственных свойств и наследует отступ первой строки от стиля
«Обычный», которому правило задаёт отступ основного текста.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_figures

DRAWING = (
    f'<w:drawing {nsdecls("w")}><wp:inline {nsdecls("wp")}>'
    f'<wp:extent cx="100" cy="100"/></wp:inline></w:drawing>'
)


@pytest.fixture()
def profile(profile_file: Path):
    return load_profile_from_path(profile_file)


def add_figure(
    document: Document,
    *,
    indent_cm: float | None = None,
    alignment=None,
):
    """Добавляет абзац с картинкой, при необходимости — с отступом автора."""
    paragraph = document.add_paragraph()
    if alignment is not None:
        paragraph.paragraph_format.alignment = alignment
    if indent_cm is not None:
        paragraph.paragraph_format.first_line_indent = Cm(indent_cm)
        paragraph.paragraph_format.left_indent = Cm(indent_cm)
    paragraph.add_run()._r.append(parse_xml(DRAWING))
    return paragraph


def indent_of(paragraph) -> dict[str, str]:
    pPr = paragraph._p.find(qn("w:pPr"))
    assert pPr is not None, "у абзаца с рисунком не заданы свойства"
    indent = pPr.find(qn("w:ind"))
    return dict(indent.attrib) if indent is not None else {}


def alignment_of(paragraph) -> str | None:
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    jc = pPr.find(qn("w:jc"))
    return None if jc is None else jc.get(qn("w:val"))


def test_figure_without_own_properties_is_centered(make_context, profile) -> None:
    """Главный случай: у абзаца вообще нет своих свойств.

    Раньше такой рисунок наследовал отступ 1,25 см от стиля «Обычный»
    и уезжал вправо.
    """
    document = Document()
    add_figure(document)

    apply_figures(make_context(document, profile))

    paragraph = document.paragraphs[0]
    assert alignment_of(paragraph) == "center"
    assert indent_of(paragraph) == {
        qn("w:firstLine"): "0",
        qn("w:left"): "0",
        qn("w:right"): "0",
    }


def test_figure_indent_from_author_is_cleared(make_context, profile) -> None:
    document = Document()
    add_figure(document, indent_cm=1.25, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)

    apply_figures(make_context(document, profile))

    paragraph = document.paragraphs[0]
    assert alignment_of(paragraph) == "center"
    assert indent_of(paragraph)[qn("w:firstLine")] == "0"
    assert indent_of(paragraph)[qn("w:left")] == "0"


def test_hanging_indent_is_removed(make_context, profile) -> None:
    """Висячий отступ и «символьные» отступы перебивают обычный."""
    document = Document()
    paragraph = add_figure(document)
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(
        parse_xml(
            f'<w:ind {nsdecls("w")} w:left="709" w:hanging="709" '
            f'w:firstLineChars="0" w:hangingChars="100"/>'
        )
    )

    apply_figures(make_context(document, profile))

    indent = indent_of(paragraph)
    assert qn("w:hanging") not in indent
    assert qn("w:hangingChars") not in indent
    assert qn("w:firstLineChars") not in indent
    assert indent[qn("w:firstLine")] == "0"


def test_figure_spacing_becomes_single(make_context, profile) -> None:
    document = Document()
    add_figure(document)

    apply_figures(make_context(document, profile))

    pPr = document.paragraphs[0]._p.find(qn("w:pPr"))
    spacing = pPr.find(qn("w:spacing"))
    assert spacing.get(qn("w:line")) == "240"
    assert spacing.get(qn("w:before")) == "0"
    assert spacing.get(qn("w:after")) == "0"


def test_title_page_figure_is_untouched(make_context, profile) -> None:
    """Рисунок на титульном листе методичка не описывает — не трогаем."""
    document = Document()
    add_figure(document, indent_cm=1.25)
    document.add_paragraph("ВВЕДЕНИЕ")

    context = make_context(document, profile)
    apply_figures(context)

    assert alignment_of(document.paragraphs[0]) is None
    assert indent_of(document.paragraphs[0])[qn("w:firstLine")] == str(Cm(1.25).twips)

    manual = context.changes.by_category()["figures"]
    assert any("титульном листе" in change.description for change in manual)


def test_alignment_comes_from_profile(make_context, profile_dict: dict, tmp_path: Path) -> None:
    data = copy.deepcopy(profile_dict)
    data["figures"]["alignment"] = "left"
    path = tmp_path / "рисунок-слева.json"
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    add_figure(document)

    apply_figures(make_context(document, load_profile_from_path(path)))

    # Даже при выравнивании по левому краю отступ убирается: он сдвигал бы рисунок.
    assert alignment_of(document.paragraphs[0]) == "left"
    assert indent_of(document.paragraphs[0])[qn("w:firstLine")] == "0"


def test_change_is_reported(make_context, profile) -> None:
    document = Document()
    add_figure(document)

    context = make_context(document, profile)
    apply_figures(context)

    changes = context.changes.by_category()["figures"]
    assert len(changes) == 1
    assert "по центру" in changes[0].description
    assert "рисунков: 1" in changes[0].detail


def test_text_paragraphs_are_not_touched(make_context, profile) -> None:
    document = Document()
    document.add_paragraph("Обычный абзац текста.")
    add_figure(document)

    apply_figures(make_context(document, profile))

    assert document.paragraphs[0]._p.find(qn("w:pPr")) is None


def test_second_run_changes_nothing(make_context, profile) -> None:
    document = Document()
    add_figure(document)

    apply_figures(make_context(document, profile))
    first = document.paragraphs[0]._p.xml

    apply_figures(make_context(document, profile))

    assert document.paragraphs[0]._p.xml == first
