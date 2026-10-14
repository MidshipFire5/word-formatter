"""Правила оформления сносок и расстановки переносов."""

from __future__ import annotations

import copy
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

from word_formatter.core.profiles import load_profile_from_path
from word_formatter.core.rules import apply_hyphenation, apply_note_styles


def ensure_style(document: Document, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def remove_style(document: Document, name: str) -> None:
    style = document.styles[name]
    style.element.getparent().remove(style.element)


def add_footnote_reference(document: Document):
    paragraph = document.add_paragraph("Текст со сноской")
    paragraph._p.append(
        parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
            '<w:footnoteReference w:id="1"/></w:r>'
        )
    )
    return paragraph


def test_footnote_style_is_formatted(make_context, profile_file: Path) -> None:
    document = Document()
    style = ensure_style(document, "Footnote Text")
    style.font.size = Pt(12)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_note_styles(context)

    assert style.font.name == "Times New Roman"
    assert style.font.size == Pt(10)
    assert style.paragraph_format.line_spacing == 1.0
    assert context.changes.count("footnotes") == 1
    assert "10 пт" in context.changes.by_category()["footnotes"][0].description


def test_missing_footnote_style_is_reported(make_context, profile_file: Path) -> None:
    document = Document()
    ensure_style(document, "Footnote Text")
    remove_style(document, "Footnote Text")
    add_footnote_reference(document)
    context = make_context(document, load_profile_from_path(profile_file))

    apply_note_styles(context)

    assert context.changes.count("footnotes") == 1
    message = context.changes.by_category()["footnotes"][0].description
    assert "стиль сносок не найден" in message


def test_document_without_notes_is_silent(make_context, profile_file: Path) -> None:
    document = Document()
    ensure_style(document, "Footnote Text")
    remove_style(document, "Footnote Text")
    document.add_paragraph("Текст без сносок.")
    context = make_context(document, load_profile_from_path(profile_file))

    apply_note_styles(context)

    assert context.changes.count("footnotes") == 0


def test_hyphenation_is_enabled(make_context, profile_file: Path) -> None:
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_hyphenation(context)

    settings = document.settings.element
    auto = settings.find(qn("w:autoHyphenation"))
    assert auto is not None
    assert auto.get(qn("w:val")) == "true"

    caps = settings.find(qn("w:doNotHyphenateCaps"))
    assert caps is not None
    assert caps.get(qn("w:val")) == "true"
    assert context.changes.count("hyphenation") == 1


def test_settings_elements_keep_schema_order(make_context, profile_file: Path) -> None:
    """Порядок элементов в w:settings задан схемой, иначе Word откажется открыть файл."""
    document = Document()
    context = make_context(document, load_profile_from_path(profile_file))

    apply_hyphenation(context)

    children = [child.tag.rsplit("}", 1)[-1] for child in document.settings.element]
    index = children.index("autoHyphenation")
    assert children[index - 1] == "defaultTabStop"
    assert children[index + 1] == "doNotHyphenateCaps"
    if "characterSpacingControl" in children:
        assert index < children.index("characterSpacingControl")


def test_disabled_hyphenation_is_reported(make_context, profile_dict: dict, tmp_path: Path) -> None:
    profile_data = copy.deepcopy(profile_dict)
    profile_data["hyphenation"]["enabled"] = False
    path = tmp_path / "без-переносов.json"
    path.write_text(json.dumps(profile_data, ensure_ascii=False), encoding="utf-8")

    document = Document()
    context = make_context(document, load_profile_from_path(path))
    apply_hyphenation(context)

    auto = document.settings.element.find(qn("w:autoHyphenation"))
    assert auto.get(qn("w:val")) == "false"
    assert document.settings.element.find(qn("w:doNotHyphenateCaps")) is None
    assert "отключена" in context.changes.by_category()["hyphenation"][0].description
