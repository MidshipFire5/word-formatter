"""Сборка примера «наивно оформленной» курсовой работы.

Документ сделан похожим на то, что пишут на самом деле: без стилей Word,
с ручными заголовками, другим шрифтом, курсивом, подчёркиванием, ручными
тире вместо списков. Такой пример удобно использовать, чтобы посмотреть,
что приложение действительно исправляет.

Запуск:  python tools/make_example.py
Файл создаётся в папке examples.
"""

from __future__ import annotations

import base64
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXAMPLES_DIR = PROJECT_ROOT / "examples"
EXAMPLE_NAME = "пример-курсовой.docx"

# Небольшая схема: 320 на 180 точек.
DIAGRAM_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAUAAAAC0CAIAAABqhmJGAAACm0lEQVR4nO3dQWrDMBBAUafkbjqzT+dCl6UtbkIif/u9bQIVEZ+ZrnTbtm0Bmj5mHwB4nIAhTMAQJmAIEzCECRjCBAxhAoYwAUPY/e+PxxjvOgnws3Vdf/nEBIYyKzScd4XeM8SBV9jzD6wJDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFDmIAhTMAQJmAIEzCECRjCLh3w+DL7FDxuXP4GLx0w1AkYwgQMYQKGMAFDmIAhTMAQJmAIEzCECRjCBAxhAoYwAUOYgCFMwBAmYAgTMIQJGMIEDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFD2H3WHz7Os4AHOcm6rkvKQX635do3aAJDmIAhbNoK3V0dT7n+PcMNLvOYwBAmYAgTMIQJGMIEDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFDmIAhTMAQJmAIEzCECRjCBAxhAoYwAUOYgCFMwBAmYAgTMIQJGMIEDGEChrD5z4tymndGeT8TGMIEDGFW6KPIvXNv5z8CExjCBAxhAoYwAUOYgCFMwBAmYAgTMIQJGMIEDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFDmIAhTMAQJmAIEzCECRjCBAxhAoYwAUOYgCFMwBA2/3lRr1TWucGJTGAIEzCETVuhcw/S840bPAITGMIEDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFDmIAhTMAQJmAIEzCECRjCBAxhAoYwAUOYgCFMwBAmYAgTMIQJGMIEDGEChjABQ5iAIUzAECZgCBMwhAkYwgQMYQKGMAFD2H3n98YYLz4J8G8mMIQJGMJu27bNPgPwIBMYwgQMYQKGMAFDmIAhTMAQJmAIEzAsXZ9frykb0C/LJwAAAABJRU5ErkJggg=="
)

BODY_TEXT = (
    "Управление персоналом представляет собой совокупность методов воздействия "
    "на коллектив с целью повышения эффективности работы организации. "
    "В современных условиях значение этого направления растёт."
)


def _add_plain(document: Document, text: str, **formatting):
    """Добавляет обычный абзац с оформлением, как это делает человек вручную."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = formatting.get("font", "Calibri")
    run.font.size = Pt(formatting.get("size", 12))
    if formatting.get("bold"):
        run.font.bold = True
    if formatting.get("italic"):
        run.font.italic = True
    if formatting.get("underline"):
        run.font.underline = True
    if formatting.get("center"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def _add_diagram(document: Document, image_path: Path) -> None:
    """Вставляет рисунок, поставив подпись **над** ним.

    Так часто делают в работах, набранных вручную: подпись сверху. Методичка
    требует обратного, поэтому на этом примере видно, как приложение
    переносит подпись под рисунок.
    """
    _add_plain(document, "Рисунок 1 – Схема управления персоналом", center=True, size=12)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(8))


def build_example(target: Path) -> Path:
    """Собирает пример документа и сохраняет его в target."""
    target.parent.mkdir(parents=True, exist_ok=True)
    image_path = target.parent / "_схема.png"
    image_path.write_bytes(DIAGRAM_PNG)

    document = Document()

    # Титульный лист: как обычно и делают — по центру, полужирным, другим шрифтом.
    for line in (
        "КРАСНОДАРСКИЙ ФИЛИАЛ ФИНУНИВЕРСИТЕТА",
        "Кафедра «Математика и информатика»",
        "",
        "КУРСОВАЯ РАБОТА",
        "по дисциплине «Управление персоналом»",
        "",
        "",
        "Выполнил: студент группы ИТ-21 Иванов И. И.",
        "Проверил: доцент Петров П. П.",
        "",
        "Краснодар, 2026",
    ):
        _add_plain(document, line, center=True, size=14 if line == "КУРСОВАЯ РАБОТА" else 12,
                   bold=line in ("КРАСНОДАРСКИЙ ФИЛИАЛ ФИНУНИВЕРСИТЕТА", "КУРСОВАЯ РАБОТА"))

    _add_plain(document, "СОДЕРЖАНИЕ", bold=True, size=12)
    _add_plain(document, "ВВЕДЕНИЕ …………………………………………………………… 3", size=12)
    _add_plain(document, "1 Теоретические основы ……………………………………… 4", size=12)

    _add_plain(document, "ВВЕДЕНИЕ", bold=True, size=12)
    _add_plain(document, "Актуальность темы обусловлена ростом конкуренции.", size=12)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Ключевое понятие — ")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    emphasised = paragraph.add_run("управление персоналом")
    emphasised.font.italic = True
    emphasised.font.size = Pt(12)
    tail = paragraph.add_run(" — рассматривается ниже.")
    tail.font.size = Pt(12)
    underlined = document.add_paragraph()
    underlined_run = underlined.add_run("Определение требует уточнения.")
    underlined_run.font.underline = True
    underlined_run.font.size = Pt(12)
    colored = document.add_paragraph()
    colored_run = colored.add_run("Этот абзац набран синим цветом и шрифтом Arial.")
    colored_run.font.name = "Arial"
    colored_run.font.size = Pt(11)
    colored_run.font.color.rgb = RGBColor(0x1F, 0x3F, 0xBF)

    _add_plain(document, "1 Теоретические основы управления персоналом", bold=True, size=12)
    _add_plain(document, "1.1 Понятие и сущность процесса", bold=True, size=12)
    for _ in range(3):
        _add_plain(document, BODY_TEXT, size=12)

    _add_plain(document, "Основные направления работы:", size=12)
    for item in (
        "– подбор и адаптация сотрудников;",
        "– обучение и развитие;",
        "– оценка результативности.",
    ):
        _add_plain(document, item, size=12)

    _add_plain(document, "Таблица 1 – Показатели деятельности организации", size=12)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    data = [
        ["Показатель", "Значение", "Единица"],
        ["Выручка", "55000", "руб."],
        ["Прочие данные", "–", "х"],
    ]
    for row_index, row in enumerate(data):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    _add_diagram(document, image_path)

    _add_plain(document, "ЗАКЛЮЧЕНИЕ", bold=True, size=12)
    _add_plain(document, "Поставленные задачи выполнены в полном объёме.", size=12)

    _add_plain(document, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", bold=True, size=12)
    _add_plain(document, "1. Никифорова Н. А. Комплексный экономический анализ. — Москва: Кнорус, 2021. — 439 с.", size=12)

    document.save(str(target))
    image_path.unlink(missing_ok=True)
    return target


def main() -> int:
    target = build_example(EXAMPLES_DIR / EXAMPLE_NAME)
    print(f"Пример создан: {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
